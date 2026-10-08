import asyncio
import hashlib
import os
import shutil
import stat
import subprocess
import sys
import unittest
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from types import SimpleNamespace

from mcp import Client
from mcp.client.stdio import StdioServerParameters
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904

import mcp_server as server


class ReadOnlyTest(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(dir=Path.cwd())
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.storage = self.root / "storage"
        self.storage.mkdir()
        self.csv = self.storage / "cargas.csv"
        self.csv.write_text("Jugador;Fecha;Minutos\nA;2026-10-01;90\nB;2026-10-02;45\nA;2026-10-03;60\n", encoding="utf-8-sig")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Eventos"
        worksheet.append(["Jugador", "Fecha", "Minutos"])
        worksheet.append(["A", datetime(2026, 10, 1), 90])
        worksheet.append(["B", datetime(2026, 10, 2), 45])
        workbook.create_sheet("Diccionario").append(["Campo", "Descripción"])
        workbook.save(self.storage / "eventos.xlsx")
        workbook.close()
        self.addCleanup(patch.stopall)
        patch.object(server, "STORAGE_DIR", self.storage).start()

    def test_missing_storage_does_not_create_it(self):
        missing = self.root / "missing"
        with patch.object(server, "STORAGE_DIR", missing):
            self.assertEqual(server.list_files(), {"files": [], "next_offset": None})
        self.assertFalse(missing.exists())

    def test_csv_filters_projection_and_pagination(self):
        first = server.read_performance_data("cargas.csv", filters={"Jugador": "A"}, columns=["Minutos"], limit=1)
        self.assertEqual(first["rows"], [{"Minutos": "90"}])
        self.assertEqual(first["next_offset"], 1)
        last = server.read_performance_data("cargas.csv", filters={"Jugador": "A"}, offset=1, limit=1)
        self.assertEqual(last["rows"][0]["Minutos"], "60")
        self.assertIsNone(last["next_offset"])
        dated = server.read_performance_data("cargas.csv", date_column="Fecha", date_from="2026-10-02", date_to="2026-10-02")
        self.assertEqual(dated["rows"][0]["Jugador"], "B")
        self.assertEqual(server.read_performance_data("cargas.csv", filters={"Jugador": "Z"})["rows"], [])

    def test_xlsx_metadata_dates_and_synthetic_dataset(self):
        metadata = server.get_file_metadata("eventos.xlsx")
        self.assertEqual(metadata["sheets"], ["Eventos", "Diccionario"])
        self.assertEqual(metadata["columns"], ["Jugador", "Fecha", "Minutos"])
        data = server.read_performance_data("eventos.xlsx", date_column="Fecha", date_from="2026-10-02")
        self.assertEqual(data["rows"], [{"Jugador": "B", "Fecha": "2026-10-02", "Minutos": 45}])
        shutil.copyfile(Path(__file__).parent / "data/defensa_performance_demo.xlsx", self.storage / "demo.xlsx")
        self.assertEqual(server.get_file_metadata("demo.xlsx")["sheets"], ["Eventos", "Jugadores", "Diccionario"])
        for sheet in ("Eventos", "Jugadores", "Diccionario"):
            self.assertTrue(server.read_performance_data("demo.xlsx", sheet=sheet, limit=1)["rows"])

    def test_excel_date_formats_and_epochs(self):
        for epoch, serial in ((CALENDAR_WINDOWS_1900, 45292), (CALENDAR_MAC_1904, 43830)):
            with self.subTest(epoch=epoch):
                workbook = Workbook()
                workbook.epoch = epoch
                worksheet = workbook.active
                worksheet.append(["Fecha", "Fecha_hora", "Numero", "Texto", "Vacio"])
                worksheet.append([serial, serial + 0.5, serial, "01/01/2024", None])
                worksheet['A2'].number_format = 'dd/mm/yyyy'
                worksheet['B2'].number_format = 'yyyy-mm-dd hh:mm:ss'
                workbook.save(self.storage / "fechas.xlsx")
                workbook.close()
                before = (self.storage / "fechas.xlsx").read_bytes()
                data = server.read_performance_data("fechas.xlsx", date_column="Fecha",
                                                    date_from="2024-01-01", date_to="2024-01-01")
                self.assertEqual(data["rows"], [{"Fecha": "2024-01-01", "Fecha_hora": "2024-01-01",
                                                "Numero": serial, "Texto": "01/01/2024", "Vacio": None}])
                self.assertEqual(server.read_performance_data("fechas.xlsx", filters={"Fecha": "2024-01-01"})["rows"], data["rows"])
                self.assertEqual(server.read_performance_data("fechas.xlsx", date_column="Fecha",
                                                             date_from="2024-01-02")["rows"], [])
                self.assertEqual((self.storage / "fechas.xlsx").read_bytes(), before)

    def test_unsafe_paths_and_unknown_inputs(self):
        for filename in ("../secret.csv", str(self.csv), "C:\\secret.csv", "storage/cargas.csv", "CON.csv", "missing.csv"):
            with self.subTest(filename=filename), self.assertRaises(ToolError):
                server.get_file_metadata(filename)
        for arguments in ({"columns": ["Missing"]}, {"columns": []}, {"filters": {"Missing": "A"}}, {"sheet": "Eventos"}, {"limit": 501}, {"offset": -1}, {"date_from": "2026-10-01"}):
            with self.subTest(arguments=arguments), self.assertRaises(ToolError):
                server.read_performance_data("cargas.csv", **arguments)
        with self.assertRaises(ToolError):
            server.get_file_metadata("eventos.xlsx", sheet="Missing")

    def test_malformed_and_unsupported_files(self):
        for name, content in (("broken.xlsx", b"invalid"), ("encoding.csv", b"\xff\xff"), ("duplicate.csv", b"A,A\n1,2"), ("empty.csv", b""), ("ragged.csv", b"A,B\n1,2,3"), ("report.pdf", b"pdf")):
            (self.storage / name).write_bytes(content)
            with self.subTest(name=name), self.assertRaises(ToolError):
                server.read_performance_data(name)
        self.assertEqual(server.get_file_metadata("report.pdf")["size_bytes"], 3)

    def test_processing_limits(self):
        with patch.object(server, "MAX_FILE_BYTES", 1), self.assertRaises(ToolError):
            server.get_file_metadata("eventos.xlsx")
        with patch.object(server, "MAX_XLSX_BYTES", 1), self.assertRaises(ToolError):
            server.read_performance_data("eventos.xlsx")
        with patch.object(server, "MAX_RESPONSE_BYTES", 20), self.assertRaises(ToolError):
            server.read_performance_data("cargas.csv")
        with patch.object(server, "MAX_ROWS", 1), self.assertRaises(ToolError):
            server.read_performance_data("cargas.csv", filters={"Jugador": "Z"})

    def test_file_link_rejected(self):
        outside = self.root / "secret.csv"
        outside.write_text("Secret\nprivate", encoding="utf-8")
        linked = self.storage / "link.csv"
        try:
            linked.symlink_to(outside)
        except OSError:
            self.skipTest("Windows requiere privilegio para crear symlinks; junction se prueba por separado.")
        with self.assertRaises(ToolError):
            server.get_file_metadata("link.csv")
        self.assertNotIn("link.csv", [item["filename"] for item in server.list_files()["files"]])

    def test_file_reparse_and_symlink_checks_without_privileges(self):
        original_lstat = Path.lstat
        for mode, attributes in ((stat.S_IFLNK, 0), (stat.S_IFREG, 0x400)):
            def lstat(path):
                if path == self.csv:
                    return SimpleNamespace(st_mode=mode, st_file_attributes=attributes)
                return original_lstat(path)
            with self.subTest(mode=mode), patch.object(Path, "lstat", lstat):
                with self.assertRaises(ToolError):
                    server.get_file_metadata("cargas.csv")
                self.assertNotIn("cargas.csv", [item["filename"] for item in server.list_files()["files"]])

    @unittest.skipUnless(os.name == "nt", "Windows junction")
    def test_storage_junction_rejected(self):
        junction = self.root / "junction"
        completed = subprocess.run(["cmd", "/c", "mklink", "/J", str(junction), str(self.storage)], capture_output=True)
        self.assertEqual(completed.returncode, 0, completed.stderr)
        try:
            with patch.object(server, "STORAGE_DIR", junction):
                with self.assertRaises(ToolError):
                    server.list_files()
                with self.assertRaises(ToolError):
                    server.get_file_metadata("cargas.csv")
        finally:
            junction.rmdir()

    def test_mcp_stdio_all_tools_and_no_writes(self):
        def snapshot():
            return {path.name: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns) for path in self.storage.iterdir()}
        before = snapshot()

        async def exercise(mode="auto"):
            parameters = StdioServerParameters(command=sys.executable, args=["-c", "from pathlib import Path; import mcp_server as s; s.STORAGE_DIR=Path(" + repr(str(self.storage)) + "); s.mcp.run(transport='stdio')"], cwd=Path(__file__).parent)
            async with Client(parameters, read_timeout_seconds=15, mode=mode) as client:
                tools = (await client.list_tools()).tools
                self.assertEqual({tool.name for tool in tools}, {"list_files", "get_file_metadata", "read_performance_data"})
                self.assertTrue(all(tool.annotations.read_only_hint for tool in tools))
                for tool, args in (("list_files", {"limit": 1}), ("get_file_metadata", {"filename": "eventos.xlsx"}), ("read_performance_data", {"filename": "cargas.csv", "filters": {"Jugador": "A"}, "limit": 1}), ("read_performance_data", {"filename": "eventos.xlsx"})):
                    result = await client.call_tool(tool, args)
                    self.assertFalse(result.is_error, result)
                    self.assertTrue(result.structured_content)
                    if tool == "read_performance_data" and args["filename"] == "eventos.xlsx":
                        self.assertEqual(result.structured_content["rows"][0]["Fecha"], "2026-10-01")
                for args in ({"filename": "../secret.csv"}, {"filename": "missing.csv"}, {"filename": "cargas.csv", "limit": 0}):
                    result = await client.call_tool("read_performance_data", args)
                    self.assertTrue(result.is_error)
                    self.assertNotIn(str(self.root), str(result.content))
                result = await client.call_tool("list_files", {"limit": "invalid"})
                self.assertTrue(result.is_error)

        async def timed():
            await asyncio.wait_for(exercise(), timeout=30)
            await asyncio.wait_for(exercise("legacy"), timeout=30)
        asyncio.run(timed())
        self.assertEqual(snapshot(), before)

    def test_stdio_script_entrypoint(self):
        async def exercise():
            parameters = StdioServerParameters(command=sys.executable, args=[str(Path(server.__file__).resolve())], cwd=self.root)
            async with Client(parameters, read_timeout_seconds=15, mode="legacy") as client:
                self.assertEqual(len((await client.list_tools()).tools), 3)
        async def timed():
            await asyncio.wait_for(exercise(), timeout=30)
        asyncio.run(timed())


if __name__ == "__main__":
    unittest.main()
